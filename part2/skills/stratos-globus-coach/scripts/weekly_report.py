#!/usr/bin/env python3
"""GLO-BUS weekly progress report (StratOS GLO-BUS Coach, Mode 5).

Usage: python weekly_report.py globus-capture-C-Y6.json globus-capture-C-Y7.json [...]
                               [--team C] [--out-dir report] [--title "Company C"]
Reads every capture file so far (stratos-globus-capture) and writes, for the latest year with results:
  weekly-report-Y<n>.html   one self-contained page with the graphs (open in a browser, print to PDF)
  weekly-report-Y<n>.docx   the same report in Word (needs python-docx; skipped if missing)
  weekly-report-Y<n>.json   the facts behind it, for the coach's narrative
Sections: the scorecard (KPIs vs investor expectations), the apparent strategic position by product,
progress graphs, where the team stands in the contest, what rivals changed (public reports only), what
to look out for next round, and questions for the team.

Rules built in: feedback, not recommendations. The report names the position the team's inputs reveal and
general things to watch; it never proposes decision values. Rival data comes only from the class-wide
reports every team can see (scoreboard_rows, cir_rows).
"""
import base64, html, io, json, os, sys

NAVY, GOLD, GREY, RED, GREEN = "#1D3557", "#C9A55C", "#8A949E", "#C44536", "#2D936C"
PRODUCTS = [("camera", "Cameras"), ("drone", "Drones")]
CREDIT = ["AAA", "AA+", "AA", "AA-", "A+", "A", "A-", "BBB+", "BBB", "BBB-", "BB+", "BB", "BB-", "B+", "B", "B-",
          "CCC+", "CCC", "CCC-", "CC", "C"]

PRINCIPLES = {
    "Low-cost provider": "Low-cost providers win by keeping cost per unit at or below the industry low and pricing "
                         "below average, while holding quality at the minimum customers will accept.",
    "Differentiation": "Differentiators win when the quality, models, warranty and image they pay for earn a price "
                       "premium large enough to cover the extra cost, and when marketing makes the difference visible.",
    "Best-cost provider": "Best-cost providers offer above-average quality at an average or slightly lower price, so "
                          "costs must stay tightly controlled to fund both.",
    "Stuck in the middle": "A position with no clear cost or quality edge usually earns middling share and margins; "
                           "teams in this spot generally decide which edge they are building before the next round.",
    "Middle of the market": "A position close to the industry average on both price and quality has no edge yet; "
                            "the strategy shows up in which way the inputs move next.",
}


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load(paths):
    caps = [json.load(open(p, encoding="utf-8")) for p in paths]
    caps.sort(key=lambda c: c.get("year", 0))
    return caps


def gap(a, b, pct=True):
    a, b = num(a), num(b)
    if a is None or b is None or (pct and b == 0):
        return None
    return a / b - 1 if pct else a - b


def classify(price_gap, pq_gap):
    if price_gap is None or pq_gap is None:
        return "Unclear", "unclear"
    if price_gap <= -0.05 and pq_gap <= 0.1:
        lab = "Low-cost provider"
        conf = "clear" if price_gap <= -0.08 else "mixed"
    elif price_gap >= 0.05 and pq_gap >= 0.2:
        lab = "Differentiation"
        conf = "clear" if price_gap >= 0.08 and pq_gap >= 0.3 else "mixed"
    elif price_gap < 0.05 and pq_gap >= 0.1:
        lab = "Best-cost provider"
        conf = "clear" if pq_gap >= 0.2 and price_gap <= 0 else "mixed"
    elif price_gap >= 0.05 and pq_gap < 0.2:
        lab = "Stuck in the middle"
        conf = "mixed"
    else:
        lab = "Middle of the market"
        conf = "mixed"
    return lab, conf


def norm_strategy(s):
    s = (s or "").lower()
    if "best" in s:
        return "Best-cost provider"
    if "low" in s:
        return "Low-cost provider"
    if "differ" in s:
        return "Differentiation"
    return None


def public_positions(cap):
    """Per company and product: average price, P/Q and share over the regions in the CIR."""
    agg = {}
    for r in cap.get("cir_rows") or []:
        agg.setdefault((str(r.get("company")), r.get("product")), []).append(r)
    out = {}
    for (co, prod), rows in agg.items():
        g = [x for x in rows if (x.get("region") or "Global").lower() == "global"] or rows
        f = lambda k: (sum(num(x.get(k)) for x in g if num(x.get(k)) is not None) /
                       max(1, sum(1 for x in g if num(x.get(k)) is not None))) if any(num(x.get(k)) is not None for x in g) else None
        out[(co, prod)] = {"price": f("price"), "pq": f("pq"), "share": f("share")}
    return out


def industry_avg(pos, prod):
    rows = [v for (co, p), v in pos.items() if p == prod and v["price"] is not None and v["pq"] is not None]
    if not rows:
        return None, None
    w = [v["share"] or 1 for v in rows]
    return (sum(v["price"] * x for v, x in zip(rows, w)) / sum(w), sum(v["pq"] * x for v, x in zip(rows, w)) / sum(w))


def group_of(price, pq, p_avg, q_avg):
    hi_p, hi_q = price > p_avg, pq >= q_avg
    return {(True, True): "Premium differentiator", (False, True): "Value leader",
            (False, False): "Low-cost / economy", (True, False): "Stuck in the middle"}[(hi_p, hi_q)]


# ---------- facts ----------
def build_facts(caps, team):
    with_res = [c for c in caps if c.get("results")]
    if not with_res:
        sys.exit("ERROR: no capture file has results yet")
    cur = with_res[-1]
    prev = with_res[-2] if len(with_res) > 1 else None
    y = cur["year"]
    R = cur["results"]
    facts = {"team": team, "year": y, "years": [c["year"] for c in with_res], "kpis": [], "products": {},
             "contest": {}, "rivals": [], "watch": [], "questions": []}
    for k, lab in (("eps", "EPS"), ("roe", "ROE"), ("stock", "Stock price"), ("credit", "Credit rating"),
                   ("image", "Image rating")):
        d = (R.get("kpis") or {}).get(k) or {}
        a, t = d.get("actual"), d.get("target")
        if k == "credit":
            met = (a in CREDIT and t in CREDIT and CREDIT.index(a) <= CREDIT.index(t)) if a and t else None
        else:
            met = (num(a) >= num(t)) if num(a) is not None and num(t) is not None else None
        show = lambda v: (f"{num(v):.1%}" if k == "roe" and num(v) is not None else
                          f"${num(v):.2f}" if k in ("eps", "stock") and num(v) is not None else v)
        facts["kpis"].append({"key": k, "label": lab, "actual": show(a), "target": show(t), "met": met})
    co = R.get("company") or {}
    facts["contest"] = {"score": co.get("score"), "rank": co.get("rank"),
                        "prev_rank": ((prev or {}).get("results") or {}).get("company", {}).get("rank") if prev else None}
    sb = {str(r.get("company")): r for r in cur.get("scoreboard_rows") or []}
    if sb:
        ranked = sorted(sb.values(), key=lambda r: num(r.get("score")) or 0, reverse=True)
        facts["contest"]["leader"] = ranked[0].get("company")
        facts["contest"]["leader_score"] = ranked[0].get("score")
        facts["contest"]["n"] = len(sb)
    stated = {p: norm_strategy((cur.get("decisions") or {}).get(f"strategy.{p}")) for p, _ in PRODUCTS}
    pos_now, pos_prev = public_positions(cur), public_positions(prev) if prev else {}
    if not pos_prev and len(caps) > 1:
        prior = [c for c in caps if c.get("year", 0) < y and c.get("cir_rows")]
        pos_prev = public_positions(prior[-1]) if prior else {}
    for p, pl in PRODUCTS:
        r = ((R.get("product") or {}).get(p)) or {}
        rp = (((prev or {}).get("results") or {}).get("product") or {}).get(p) or {}
        pg, qg, cg = gap(r.get("price"), r.get("ind_price")), gap(r.get("pq"), r.get("ind_pq"), pct=False), \
            gap(r.get("cost_unit"), r.get("ind_cost_unit"))
        pg0, qg0, cg0 = gap(rp.get("price"), rp.get("ind_price")), gap(rp.get("pq"), rp.get("ind_pq"), pct=False), \
            gap(rp.get("cost_unit"), rp.get("ind_cost_unit"))
        lab, conf = classify(pg, qg)
        shares = [v for v in (r.get("share") or {}).values() if num(v) is not None]
        shares0 = [v for v in (rp.get("share") or {}).values() if num(v) is not None]
        share = sum(shares) / len(shares) if shares else None
        share0 = sum(shares0) / len(shares0) if shares0 else None
        focus = None
        if r.get("share") and shares and max(shares) - min(shares) >= 3:
            names = {"na": "North America", "ea": "Europe-Africa", "ap": "Asia-Pacific", "la": "Latin America"}
            top = max(r["share"], key=lambda k: num(r["share"][k]) or 0)
            focus = names.get(top, top)
        facts["products"][p] = {"label": pl, "apparent": lab, "confidence": conf, "stated": stated[p],
                                "price_gap": pg, "pq_gap": qg, "cost_gap": cg, "price_gap_prev": pg0,
                                "pq_gap_prev": qg0, "cost_gap_prev": cg0, "share": share, "share_prev": share0,
                                "op_margin": num(r.get("op_margin")), "op_margin_prev": num(rp.get("op_margin")),
                                "strongest_region": focus}
        # rivals: public positions only
        pa, qa = industry_avg(pos_now, p)
        mine = pos_now.get((team, p))
        if pa and mine and mine["price"] is not None:
            my_group = group_of(mine["price"], mine["pq"], pa, qa)
            facts["products"][p]["group"] = my_group
            for (c2, p2), v in pos_now.items():
                if p2 != p or c2 == team or v["price"] is None:
                    continue
                g_now = group_of(v["price"], v["pq"], pa, qa)
                old = pos_prev.get((c2, p))
                pa0, qa0 = industry_avg(pos_prev, p) if pos_prev else (None, None)
                g_old = group_of(old["price"], old["pq"], pa0, qa0) if old and pa0 else None
                facts["rivals"].append({"company": c2, "product": p, "group": g_now, "group_prev": g_old,
                                        "price_change": gap(v["price"], (old or {}).get("price")),
                                        "pq_change": gap(v["pq"], (old or {}).get("pq"), pct=False),
                                        "share": v["share"], "share_change": gap(v["share"], (old or {}).get("share"), pct=False),
                                        "entered_your_group": g_now == my_group and g_old is not None and g_old != my_group,
                                        "score": num((sb.get(c2) or {}).get("score"))})
    watch(facts, cur, prev)
    return facts


def watch(f, cur, prev):
    W, Q = f["watch"], f["questions"]
    add = lambda title, why, lesson: W.append({"title": title, "evidence": why, "lesson": lesson})
    for p, d in f["products"].items():
        pl = d["label"]
        if d["stated"] and d["apparent"] not in ("Unclear",) and d["stated"] != d["apparent"]:
            add(f"{pl}: the position your inputs show differs from the strategy you chose",
                f"You chose {d['stated'].lower()}; price {fmt_pct(d['price_gap'])} vs the industry and P/Q "
                f"{fmt_pts(d['pq_gap'])} stars put you closer to {d['apparent'].lower()}.",
                "Teams that drift between positions usually pay for both and get credit for neither. "
                "Decide whether the drift is deliberate.")
            Q.append(f"Is the {pl.lower()} position your inputs now show the one you meant to build?")
        if d["apparent"] == "Low-cost provider" and d["cost_gap"] is not None and d["cost_gap"] > -0.02:
            add(f"{pl}: the cost advantage is thin",
                f"Cost per unit is {fmt_pct(d['cost_gap'])} vs the industry while price is {fmt_pct(d['price_gap'])}.",
                PRINCIPLES["Low-cost provider"])
        if d["apparent"] in ("Differentiation", "Best-cost provider") and d["pq_gap"] is not None and \
                d["pq_gap_prev"] is not None and d["pq_gap"] < d["pq_gap_prev"] - 0.05:
            add(f"{pl}: the quality edge is narrowing",
                f"P/Q lead over the industry went from {fmt_pts(d['pq_gap_prev'])} to {fmt_pts(d['pq_gap'])} stars.",
                PRINCIPLES[d["apparent"]])
        if d["cost_gap"] is not None and d["cost_gap_prev"] is not None and d["cost_gap"] > d["cost_gap_prev"] + 0.02:
            add(f"{pl}: cost per unit is rising relative to the industry",
                f"From {fmt_pct(d['cost_gap_prev'])} to {fmt_pct(d['cost_gap'])} vs the industry average.",
                "Whatever the strategy, a widening cost gap has to be earned back in price or volume.")
        if d["share"] is not None and d["share_prev"] is not None and d["share"] < d["share_prev"] - 0.3:
            add(f"{pl}: market share slipped",
                f"Average share across regions fell from {d['share_prev']:.1f}% to {d['share']:.1f}%.",
                "Share usually moves because rivals changed price, P/Q or marketing more than you did; "
                "the CIR shows which.")
        if d["op_margin"] is not None and d["op_margin_prev"] is not None and d["op_margin"] < d["op_margin_prev"] - 0.01:
            add(f"{pl}: operating margin fell",
                f"From {d['op_margin_prev']:.1%} to {d['op_margin']:.1%}.",
                "A falling margin with flat or rising share often means the volume was bought; check which costs grew.")
    for k in f["kpis"]:
        if k["met"] is False and k["key"] in ("eps", "image", "credit"):
            add(f"{k['label']} below investor expectations",
                f"{k['label']}: {k['actual']} against an expectation of {k['target']}.",
                {"eps": "EPS is scored directly; it reflects price, cost and volume together, so look for which of "
                        "the three moved against you.",
                 "image": "Image rating builds slowly from P/Q, models, warranty, advertising and CSR; it is hard "
                          "to recover quickly once it falls behind.",
                 "credit": "Credit rating follows debt, interest coverage and cash; it also affects borrowing costs."}[k["key"]])
    c = f["contest"]
    if c.get("rank") and c.get("prev_rank") and c["rank"] > c["prev_rank"]:
        add("You lost ground in the standings", f"Rank went from {c['prev_rank']} to {c['rank']}.",
            "Rank is relative: rivals' moves matter as much as yours. The Rivals section shows who moved.")
    for r in f["rivals"]:
        if r["entered_your_group"]:
            add(f"Company {r['company']} moved into your {r['product']} group",
                f"It is now in the {r['group'].lower()} group with you (price {fmt_pct(r['price_change'])}, "
                f"P/Q {fmt_pts(r['pq_change'])} stars vs last year).",
                "When a rival moves onto your ground, price pressure in that group usually rises. Watch its "
                "next move and what it does to your share.")
    if not Q:
        Q.append("Which one result from this year surprised the team most, and what does it say about your assumptions?")
    Q.append("Which input did you change most this year, and did the result move the way the projections said?")
    if any(w["title"].endswith("thin") or "cost" in w["title"] for w in W):
        Q.append("Which cost lines grew fastest this year, and are they ones customers notice?")
    f["questions"] = Q[:3]


def fmt_pct(x):
    return "n/a" if x is None else f"{x:+.1%}"


def fmt_pts(x):
    return "n/a" if x is None else f"{x:+.1f}"


# ---------- charts ----------
def charts(caps, facts, team):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.edgecolor": "#C8CDD3", "axes.titleweight": "bold",
                         "axes.titlecolor": NAVY, "axes.titlesize": 10})
    res = [c for c in caps if c.get("results")]
    ys = [c["year"] for c in res]
    out = {}

    def save(fig, name):
        b = io.BytesIO()
        fig.savefig(b, format="png", dpi=150, bbox_inches="tight")
        plt.close(fig)
        out[name] = b.getvalue()

    def series(path):
        v = []
        for c in res:
            d = c["results"]
            for k in path:
                d = (d or {}).get(k) if isinstance(d, dict) else None
            v.append(num(d))
        return v

    # 1 KPIs vs expectations
    fig, axs = plt.subplots(1, 4, figsize=(11, 2.4))
    for ax, (k, lab, f_) in zip(axs, [("eps", "EPS ($)", "{:.2f}"), ("roe", "ROE", "{:.0%}"),
                                      ("stock", "Stock price ($)", "{:.0f}"), ("image", "Image rating", "{:.0f}")]):
        a, t = series(["kpis", k, "actual"]), series(["kpis", k, "target"])
        ax.plot(ys, a, color=NAVY, lw=2.2, marker="o", label="You")
        ax.plot(ys, t, color=GOLD, lw=1.6, ls="--", marker="o", ms=3, label="Investor expectation")
        ax.set_title(lab)
        ax.set_xticks(ys)
        ax.set_xticklabels([f"Y{y}" for y in ys])
        if k == "roe":
            ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0, decimals=1))
    axs[0].legend(frameon=False, fontsize=7, loc="upper left")
    save(fig, "kpis")

    # 2 position vs industry: price, P/Q, cost per unit
    fig, axs = plt.subplots(1, 3, figsize=(11, 2.6))
    for p, pl in PRODUCTS:
        col = NAVY if p == "camera" else GOLD
        pr = series(["product", p, "price"]); ip = series(["product", p, "ind_price"])
        q = series(["product", p, "pq"]); iq = series(["product", p, "ind_pq"])
        cu = series(["product", p, "cost_unit"]); ic = series(["product", p, "ind_cost_unit"])
        axs[0].plot(ys, [a / b - 1 if a and b else None for a, b in zip(pr, ip)], color=col, lw=2, marker="o", label=pl)
        axs[1].plot(ys, [a - b if a is not None and b is not None else None for a, b in zip(q, iq)], color=col, lw=2, marker="o", label=pl)
        axs[2].plot(ys, [a / b - 1 if a and b else None for a, b in zip(cu, ic)], color=col, lw=2, marker="o", label=pl)
    for ax, t in zip(axs, ["Price vs industry average", "P/Q vs industry (stars)", "Cost per unit vs industry"]):
        ax.axhline(0, color=GREY, lw=0.8)
        ax.set_title(t)
        ax.set_xticks(ys)
        ax.set_xticklabels([f"Y{y}" for y in ys])
    for ax in (axs[0], axs[2]):
        ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0, decimals=0))
    axs[0].legend(frameon=False, fontsize=7)
    save(fig, "position_trend")

    # 3 share and margin
    fig, axs = plt.subplots(1, 2, figsize=(11, 2.4))
    for p, pl in PRODUCTS:
        col = NAVY if p == "camera" else GOLD
        sh = []
        for c in res:
            s = ((c["results"].get("product") or {}).get(p) or {}).get("share") or {}
            v = [num(x) for x in s.values() if num(x) is not None]
            sh.append(sum(v) / len(v) if v else None)
        axs[0].plot(ys, sh, color=col, lw=2, marker="o", label=pl)
        axs[1].plot(ys, series(["product", p, "op_margin"]), color=col, lw=2, marker="o", label=pl)
    axs[0].set_title("Market share (average of regions, %)")
    axs[1].set_title("Operating margin")
    axs[1].yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0, decimals=0))
    for ax in axs:
        ax.set_xticks(ys)
        ax.set_xticklabels([f"Y{y}" for y in ys])
    axs[0].legend(frameon=False, fontsize=7)
    save(fig, "share_margin")

    # 4 contest: scores by company
    sb_years = [c for c in caps if c.get("scoreboard_rows")]
    if sb_years:
        fig, ax = plt.subplots(figsize=(11, 2.8))
        comps = sorted({str(r["company"]) for c in sb_years for r in c["scoreboard_rows"]})
        yy = [c["year"] for c in sb_years]
        for co in comps:
            v = [num(next((r.get("score") for r in c["scoreboard_rows"] if str(r["company"]) == co), None)) for c in sb_years]
            mine = co == team
            ax.plot(yy, v, color=GOLD if mine else "#B8C2CC", lw=3 if mine else 1.2, marker="o", ms=5 if mine else 3,
                    zorder=3 if mine else 1)
            if v[-1] is not None:
                ax.annotate(f"{co}{' (you)' if mine else ''}", (yy[-1], v[-1]), xytext=(6, 0), textcoords="offset points",
                            va="center", fontsize=8, color=NAVY if mine else GREY, fontweight="bold" if mine else "normal")
        ax.set_title("Overall score by company (class scoreboard)")
        ax.set_xticks(yy)
        ax.set_xticklabels([f"Y{y}" for y in yy])
        ax.set_xlim(yy[0] - 0.2, yy[-1] + 0.5)
        save(fig, "contest")

    # 5 position maps from the CIR (latest year, arrows from the year before)
    cir_caps = [c for c in caps if c.get("cir_rows")]
    if cir_caps:
        now = public_positions(cir_caps[-1])
        before = public_positions(cir_caps[-2]) if len(cir_caps) > 1 else {}
        fig, axs = plt.subplots(1, 2, figsize=(11, 3.6))
        for ax, (p, pl) in zip(axs, PRODUCTS):
            pa, qa = industry_avg(now, p)
            if pa is None:
                ax.set_visible(False)
                continue
            for (co, p2), v in now.items():
                if p2 != p or v["price"] is None:
                    continue
                mine = co == team
                old = before.get((co, p))
                if old and old["price"] is not None:
                    ax.annotate("", xy=(v["pq"], v["price"]), xytext=(old["pq"], old["price"]),
                                arrowprops=dict(arrowstyle="->", color=GOLD if mine else "#C8CDD3", lw=1.4 if mine else 0.8))
                ax.scatter(v["pq"], v["price"], s=40 + 18 * (v["share"] or 10), color=GOLD if mine else NAVY,
                           alpha=0.95 if mine else 0.55, edgecolor="white", zorder=3)
                ax.annotate(co, (v["pq"], v["price"]), ha="center", va="center", fontsize=8, color="white",
                            fontweight="bold", zorder=4)
            ax.axvline(qa, color=GREY, lw=0.8, ls=":")
            ax.axhline(pa, color=GREY, lw=0.8, ls=":")
            ax.set_xlabel("P/Q rating (stars)")
            ax.set_ylabel("Price ($)")
            ax.set_title(f"{pl}: positions in Year {cir_caps[-1]['year']} (arrows from Year {cir_caps[-2]['year']})"
                         if before else f"{pl}: positions in Year {cir_caps[-1]['year']}")
            for txt, xa, ya, ha, va in (("Premium", 0.98, 0.98, "right", "top"), ("Value", 0.98, 0.02, "right", "bottom"),
                                        ("Economy", 0.02, 0.02, "left", "bottom"), ("Overpriced", 0.02, 0.98, "left", "top")):
                ax.text(xa, ya, txt, transform=ax.transAxes, ha=ha, va=va, fontsize=7, color=GREY)
        save(fig, "maps")
    return out


# ---------- report ----------
def sections(f):
    """Text blocks shared by the HTML and Word versions."""
    c = f["contest"]
    snap = []
    if c.get("rank"):
        s = f"Rank {c['rank']}" + (f" of {c['n']}" if c.get("n") else "") + f", overall score {c.get('score')}"
        if c.get("prev_rank"):
            s += f" (rank {c['prev_rank']} last year)"
        if c.get("leader") and c.get("leader") != f["team"]:
            s += f". Leader: Company {c['leader']} at {c.get('leader_score')}."
        snap.append(s)
    pos = []
    for p, d in f["products"].items():
        t = (f"{d['label']}: your inputs look like a **{d['apparent'].lower()}** position ({d['confidence']}). "
             f"Price {fmt_pct(d['price_gap'])} vs the industry, P/Q {fmt_pts(d['pq_gap'])} stars, cost per unit "
             f"{fmt_pct(d['cost_gap'])}.")
        if d.get("group"):
            t += f" On the competitor report you sit in the {d['group'].lower()} group."
        if d.get("strongest_region"):
            t += f" Your strongest region is {d['strongest_region']}."
        if d["stated"]:
            t += (" That matches the strategy you chose." if d["stated"] == d["apparent"]
                  else f" You chose {d['stated'].lower()}.")
        t += " " + PRINCIPLES.get(d["apparent"], "")
        pos.append(t)
    riv = []
    for r in sorted(f["rivals"], key=lambda r: -(abs(r["price_change"] or 0) + abs(r["pq_change"] or 0) / 5))[:5]:
        if r["price_change"] is None:
            continue
        riv.append(f"Company {r['company']} ({r['product']}s): price {fmt_pct(r['price_change'])}, P/Q "
                   f"{fmt_pts(r['pq_change'])} stars, share {fmt_pts(r['share_change'])} points; now in the "
                   f"{r['group'].lower()} group" + (" (moved into yours)" if r["entered_your_group"] else "") + ".")
    return snap, pos, riv


def md_bold(t):
    import re
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", html.escape(t))


def write_html(f, imgs, path, title):
    snap, pos, riv = sections(f)
    img = lambda k: (f'<img src="data:image/png;base64,{base64.b64encode(imgs[k]).decode()}" alt="{k}">' if k in imgs else "")
    krows = "".join(
        f"<tr><td>{html.escape(k['label'])}</td><td>{html.escape(str(k['actual']))}</td><td>{html.escape(str(k['target']))}</td>"
        f"<td class='{'ok' if k['met'] else ('no' if k['met'] is False else '')}'>"
        f"{'Met' if k['met'] else ('Below' if k['met'] is False else 'n/a')}</td></tr>" for k in f["kpis"])
    watch = "".join(f"<li><b>{html.escape(w['title'])}.</b> {html.escape(w['evidence'])} <i>{html.escape(w['lesson'])}</i></li>"
                    for w in f["watch"]) or "<li>Nothing stands out against your position this round.</li>"
    qs = "".join(f"<li>{html.escape(q)}</li>" for q in f["questions"])
    page = f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>body{{font-family:Arial,sans-serif;color:#1F2933;max-width:980px;margin:24px auto;padding:0 16px;line-height:1.45}}
h1{{font-family:Cambria,Georgia,serif;color:{NAVY};margin-bottom:2px}}h2{{font-family:Cambria,Georgia,serif;color:{NAVY};
border-bottom:2px solid {GOLD};padding-bottom:3px;margin-top:28px}}.sub{{color:#B5651D;font-weight:bold}}
table{{border-collapse:collapse}}td,th{{border:1px solid #C8CDD3;padding:4px 10px}}th{{background:{NAVY};color:#fff}}
.ok{{background:#D8F0DF}}.no{{background:#F8D7D3}}img{{max-width:100%}}.note{{color:#5F6B76;font-size:12px}}
li{{margin-bottom:6px}}</style></head><body>
<h1>{html.escape(title)}</h1><div class="sub">GLO-BUS weekly report · Year {f['year']} results · StratOS</div>
<p class="note">Feedback, not recommendations: this report names the position your inputs reveal and general things to
watch. Every decision is your team's. Rival figures come only from the class-wide reports every team can see.</p>
<h2>1. Your scorecard</h2>{''.join(f'<p>{html.escape(s)}</p>' for s in snap)}
<table><tr><th>Scored measure</th><th>You</th><th>Investor expectation</th><th></th></tr>{krows}</table>
{img('kpis')}
<h2>2. The position you've taken</h2>{''.join(f'<p>{md_bold(t)}</p>' for t in pos)}{img('maps')}{img('position_trend')}
<h2>3. Progress</h2>{img('share_margin')}
<h2>4. Where you stand in the contest</h2>{img('contest')}
{'<p><b>What rivals changed since last year</b> (from the competitor report):</p><ul>' + ''.join(f'<li>{html.escape(r)}</li>' for r in riv) + '</ul>' if riv else ''}
<h2>5. What to look out for next round</h2><ul>{watch}</ul>
<h2>6. Questions for your team</h2><ol>{qs}</ol>
<p class="note">StratOS Strategy Lab · ©2026 G. Bradley Scheller · Generated by weekly_report.py</p></body></html>"""
    open(path, "w", encoding="utf-8").write(page)


def write_docx(f, imgs, path, title):
    try:
        import docx
        from docx.shared import Pt, RGBColor, Inches
    except ImportError:
        return False
    snap, pos, riv = sections(f)
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.8)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(10.5)
    navy = RGBColor(0x1D, 0x35, 0x57)

    def h(t, size=14):
        p = d.add_paragraph()
        r = p.add_run(t)
        r.bold, r.font.size, r.font.name, r.font.color.rgb = True, Pt(size), "Cambria", navy
        return p

    def para(t):
        p = d.add_paragraph()
        for i, part in enumerate(t.split("**")):
            p.add_run(part).bold = i % 2 == 1
        return p

    def pic(k):
        if k in imgs:
            d.add_picture(io.BytesIO(imgs[k]), width=Inches(6.9))

    h(title, 20)
    p = d.add_paragraph()
    r = p.add_run(f"GLO-BUS weekly report · Year {f['year']} results · StratOS")
    r.bold, r.font.color.rgb = True, RGBColor(0xB5, 0x65, 0x1D)
    p = d.add_paragraph()
    r = p.add_run("Feedback, not recommendations: this report names the position your inputs reveal and general things "
                  "to watch. Every decision is your team's. Rival figures come only from the class-wide reports.")
    r.italic, r.font.size = True, Pt(9)
    h("1. Your scorecard")
    for s in snap:
        para(s)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, v in zip(t.rows[0].cells, ["Scored measure", "You", "Investor expectation", ""]):
        c.text = v
    for k in f["kpis"]:
        row = t.add_row().cells
        for c, v in zip(row, [k["label"], str(k["actual"]), str(k["target"]),
                              "Met" if k["met"] else ("Below" if k["met"] is False else "n/a")]):
            c.text = v
    pic("kpis")
    h("2. The position you've taken")
    for s in pos:
        para(s)
    pic("maps")
    pic("position_trend")
    h("3. Progress")
    pic("share_margin")
    h("4. Where you stand in the contest")
    pic("contest")
    if riv:
        para("**What rivals changed since last year** (from the competitor report):")
        for s in riv:
            d.add_paragraph(s, style="List Bullet")
    h("5. What to look out for next round")
    for w in f["watch"] or [{"title": "Nothing stands out against your position this round", "evidence": "", "lesson": ""}]:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(w["title"] + ". ").bold = True
        p.add_run(w["evidence"] + " ")
        p.add_run(w["lesson"]).italic = True
    h("6. Questions for your team")
    for q in f["questions"]:
        d.add_paragraph(q, style="List Number")
    d.save(path)
    return True


def main():
    a = sys.argv[1:]
    if not a:
        sys.exit(__doc__)
    opt = lambda k, dflt=None: a[a.index(k) + 1] if k in a else dflt
    files = [x for i, x in enumerate(a) if x.endswith(".json") and (i == 0 or a[i - 1] not in ("--team", "--out-dir", "--title"))]
    caps = load(files)
    team = opt("--team") or next((str(c.get("company")) for c in caps if c.get("company")), None)
    out = opt("--out-dir", ".")
    os.makedirs(out, exist_ok=True)
    facts = build_facts(caps, team)
    y = facts["year"]
    title = opt("--title", f"Company {team}")
    imgs = charts(caps, facts, team)
    base = os.path.join(out, f"weekly-report-Y{y}")
    write_html(facts, imgs, base + ".html", title)
    ok = write_docx(facts, imgs, base + ".docx", title)
    json.dump(facts, open(base + ".json", "w"), indent=1, default=lambda v: round(v, 4) if isinstance(v, float) else str(v))
    print(json.dumps({"year": y, "html": base + ".html", "docx": base + ".docx" if ok else None, "json": base + ".json",
                      "apparent": {p: d["apparent"] for p, d in facts["products"].items()},
                      "watch_items": len(facts["watch"])}))


if __name__ == "__main__":
    main()
