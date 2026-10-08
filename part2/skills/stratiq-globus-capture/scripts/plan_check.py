#!/usr/bin/env python3
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
# Copyright (c) 2026 G. Bradley Scheller · Strat-IQ Toolkit · commercial licences: BScheller@ToolsIQ.ai
"""Compare the decisions a team entered in GLO-BUS (from a capture file) with its plan in the Strat-IQ
Workbook's GLO-BUS Planner sheet. Read-only: it reports differences, it never changes anything.

Usage: python plan_check.py StratIQ_Workbook.xlsx globus-capture-C-Y7.json [--year 7] [--tolerance 0.005]
  The capture file must hold a "decisions" object keyed like the planner's Key column
  (e.g. "mkt.camera.na.price": 249). --year defaults to the capture's "year".
  Numbers match when they differ by no more than the tolerance (fraction, default 0.5%) or 0.01.
  Text matches ignoring case and surrounding spaces.
"""
import json, sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("ERROR: openpyxl is required (pip install openpyxl)")

SHEET = "GLO-BUS Planner"


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        t = v.strip().replace(",", "").replace("$", "").replace("%", "")
        try:
            return float(t)
        except ValueError:
            return None
    return None


def load_plan(path, year):
    wb = load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"ERROR: no '{SHEET}' sheet in {path}")
    ws = wb[SHEET]
    col = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(row=5, column=c).value or "").strip().lower() == f"year {year}":
            col = c
            break
    if col is None:
        sys.exit(f"ERROR: no 'Year {year}' column in the planner")
    plan, labels = {}, {}
    for r in range(6, ws.max_row + 1):
        key = ws.cell(row=r, column=1).value
        if not key or not isinstance(key, str) or "." not in key or key.lower() == "key":
            continue
        if key in plan:          # the change block repeats the keys further down; keep the first (inputs)
            continue
        plan[key] = ws.cell(row=r, column=col).value
        labels[key] = ws.cell(row=r, column=3).value or key
    return plan, labels


def same(a, b, tol):
    na, nb = num(a), num(b)
    if na is not None and nb is not None:
        return abs(na - nb) <= max(0.01, tol * max(abs(na), abs(nb)))
    return str(a).strip().lower() == str(b).strip().lower()


def main():
    a = sys.argv
    if len(a) < 3:
        sys.exit(__doc__)
    cap = json.load(open(a[2], encoding="utf-8"))
    year = int(a[a.index("--year") + 1]) if "--year" in a else cap.get("year")
    if year is None:
        sys.exit("ERROR: give --year or a capture with a 'year'")
    tol = float(a[a.index("--tolerance") + 1]) if "--tolerance" in a else 0.005
    plan, labels = load_plan(a[1], year)
    entered = cap.get("decisions") or {}
    if not entered:
        sys.exit("ERROR: the capture has no 'decisions' values; capture the decision screens first")
    mismatches, not_seen, unplanned, matched = [], [], [], 0
    for k, pv in plan.items():
        if pv in (None, ""):
            continue
        if k not in entered or entered[k] in (None, ""):
            not_seen.append({"key": k, "decision": labels[k], "planned": pv})
        elif same(pv, entered[k], tol):
            matched += 1
        else:
            mismatches.append({"key": k, "decision": labels[k], "planned": pv, "entered": entered[k]})
    for k, ev in entered.items():
        if ev not in (None, "") and plan.get(k) in (None, ""):
            unplanned.append({"key": k, "decision": labels.get(k, k), "entered": ev})
    out = {"year": year, "company": cap.get("company"), "matched": matched, "mismatches": mismatches,
           "planned_not_seen_on_screens": not_seen, "entered_but_not_planned": unplanned,
           "all_match": not mismatches and not not_seen}
    txt = json.dumps(out, indent=2, default=str)
    if "--out" in a:
        open(a[a.index("--out") + 1], "w").write(txt)
    print(txt)


if __name__ == "__main__":
    main()
