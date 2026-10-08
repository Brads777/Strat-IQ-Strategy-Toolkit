#!/usr/bin/env python3
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
# Copyright (c) 2026 G. Bradley Scheller · Strat-IQ Toolkit · commercial licences: BScheller@ToolsIQ.ai
"""Read a Strat-IQ workbook back into the strategy ledger (the workbook's two-way sync).

build_workbook.py records, in a hidden sheet called _ledger_map, which ledger field every blue-on-cream input
cell holds and the value it had when the workbook was built. This script compares each input cell with that
value and writes only the cells the team changed back into the ledger. Fields the workbook does not show
(evidence lists, notes, ids it displays as names) are left exactly as they were.

  python read_workbook.py StratIQ_Workbook.xlsx --ledger strategy-ledger.json            # update in place (backup kept)
  python read_workbook.py StratIQ_Workbook.xlsx --ledger strategy-ledger.json --dry-run  # list changes only
  python read_workbook.py StratIQ_Workbook.xlsx --out strategy-ledger.json               # blank workbook -> new ledger

Rules: a row whose input cells are all cleared removes that item; a new row adds one; edit rows in place
(do not sort or move rows, because rows map to list positions). Nothing here enters anything in GLO-BUS.
"""
import argparse, datetime, json, os, re, shutil, sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter

PESTEL = {"political": "P", "economic": "E", "social": "S", "technological": "T", "environmental": "En", "legal": "L"}
PERSP = {"financial": "financial", "customer": "customer", "internal process": "internal", "internal": "internal",
         "learning and growth": "learning", "learning": "learning"}


def norm(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return round(float(v), 9)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return str(v)[:10]
    s = str(v).strip()
    if not s:
        return None
    try:
        return round(float(s), 9) if re.fullmatch(r"-?\d+(\.\d+)?", s) else s
    except ValueError:
        return s


def placeholder(v):
    return isinstance(v, str) and v.strip().startswith("[") and v.strip().endswith("]")


def convert(v, conv):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        if not v:
            return None
    if isinstance(v, (datetime.date, datetime.datetime)):
        v = str(v)[:10]
    if isinstance(v, float) and v.is_integer() and conv in ("int", "year"):
        v = int(v)
    if conv == "list":
        return [x.strip() for x in re.split(r"[;,]", str(v)) if x.strip()]
    if conv == "pairs":
        return [x.strip() for x in re.split(r"\s*[×x,]\s*", str(v)) if x.strip()]
    if conv == "bool":
        return {"yes": True, "no": False}.get(str(v).strip().lower(), v)
    if conv == "yn":
        return {"yes": "yes", "no": "no", "?": "?"}.get(str(v).strip().lower(), v)
    if conv == "div10":
        return v / 10 if isinstance(v, (int, float)) else v
    if conv == "pestel":
        return PESTEL.get(str(v).strip().lower(), v)
    if conv == "persp":
        return PERSP.get(str(v).strip().lower(), v)
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


TOKEN = re.compile(r"^(.*?)(?:\[([^\]]*)\])?$")


class Applier:
    def __init__(self, root):
        self.root = root
        self.new_rows = {}     # (sheet, row, id(list)) -> item created by '[+]'
        self.padded = []       # (list, item) pads created to reach an index

    def walk(self, toks, create, rowkey):
        """Return (container, key) for the last token, creating the path when create is True."""
        cur = self.root
        for n, tok in enumerate(toks):
            name, sel = TOKEN.match(tok).groups()
            last = n == len(toks) - 1
            if not isinstance(cur, dict):
                return None
            if sel is None:
                if last:
                    return cur, name
                nxt = cur.get(name)
                if not isinstance(nxt, (dict, list)):
                    if not create:
                        return None
                    nxt = cur[name] = {}
                cur = nxt
                continue
            lst = cur.get(name)
            if not isinstance(lst, list):
                if not create:
                    return None
                lst = cur[name] = []
            if sel.isdigit():
                i = int(sel)
                if last:
                    while len(lst) <= i:
                        if not create:
                            return None
                        lst.append(None)
                    return lst, i
                while len(lst) <= i:
                    if not create:
                        return None
                    pad = {}
                    lst.append(pad)
                    self.padded.append((lst, pad))
                item = lst[i]
            elif sel == "+":
                k = (rowkey, id(lst))
                item = self.new_rows.get(k)
                if item is None:
                    if not create:
                        return None
                    item = {}
                    lst.append(item)
                    self.new_rows[k] = item
            else:
                fk, fv = sel.split("=", 1)
                item = next((x for x in lst if isinstance(x, dict) and str(x.get(fk)) == fv), None)
                if item is None:
                    if not create:
                        return None
                    item = {fk: fv}
                    lst.append(item)
            if isinstance(item, str):  # e.g. a takeaway stored as plain text
                idx = lst.index(item)
                item = lst[idx] = {"text": item}
            if item is None:
                if not create:
                    return None
                idx = lst.index(None)
                item = lst[idx] = {}
            if last:
                return None
            cur = item
        return None

    def get(self, toks):
        loc = self.walk(toks, False, None)
        if not loc:
            return None
        c, k = loc
        try:
            return c[k] if isinstance(c, list) else c.get(k)
        except (IndexError, KeyError):
            return None

    def set(self, toks, value, rowkey):
        loc = self.walk(toks, value is not None, rowkey)
        if not loc:
            return False
        c, k = loc
        if isinstance(c, list):
            c[k] = value
        elif value is None:
            c.pop(k, None) if k in c else None
        else:
            c[k] = value
        return True

    def item_of(self, toks):
        """The list and the item a path ends inside (the last [..] selector), for removing cleared rows."""
        last_sel = max((i for i, t in enumerate(toks) if TOKEN.match(t).group(2) not in (None, "+")), default=None)
        if last_sel is None or last_sel == len(toks) - 1:
            return None
        parent = self.walk(toks[:last_sel] + ["__probe__"], False, None) if last_sel else (self.root, None)
        name, sel = TOKEN.match(toks[last_sel]).groups()
        cont = parent[0] if parent else None
        lst = cont.get(name) if isinstance(cont, dict) else None
        if not isinstance(lst, list):
            return None
        if sel.isdigit():
            i = int(sel)
            return (lst, lst[i]) if i < len(lst) and isinstance(lst[i], dict) else None
        fk, fv = sel.split("=", 1)
        it = next((x for x in lst if isinstance(x, dict) and str(x.get(fk)) == fv), None)
        return (lst, it) if it is not None else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--ledger", help="ledger to update (omit for a workbook built --blank or as an example)")
    ap.add_argument("--out", help="write here instead of updating --ledger in place")
    ap.add_argument("--dry-run", action="store_true", help="list the changes; write nothing")
    ap.add_argument("--force", action="store_true", help="apply even if the workbook was built from another ledger")
    a = ap.parse_args()

    wb = load_workbook(a.workbook)                     # input values (and formulas if someone typed one)
    wbv = load_workbook(a.workbook, data_only=True)    # Excel's cached results for typed formulas
    if "_ledger_map" not in wb.sheetnames:
        sys.exit("This workbook has no _ledger_map sheet. Rebuild it with build_workbook.py (v3.1 or later), copy your "
                 "inputs across, then read it back.")
    mp = wb["_ledger_map"]
    rows = list(mp.iter_rows(values_only=True))
    meta, binds = rows[0], [r for r in rows[2:] if r and r[0]]
    L = json.load(open(a.ledger, encoding="utf-8")) if a.ledger else {}
    if a.ledger and meta[3] and L.get("scope_id") and meta[3] != L.get("scope_id") and not a.force:
        sys.exit(f"The workbook was built from ledger scope '{meta[3]}' but this ledger is '{L.get('scope_id')}'. "
                 "Use --force if that is intended.")

    def cell(sheet, ref):
        v = wb[sheet][ref].value
        if isinstance(v, str) and v.startswith("="):
            v = wbv[sheet][ref].value
        return v

    def resolve(path, sheet, ref):
        row = coordinate_to_tuple(ref)[0]
        out = []
        for tok in path.split("/"):
            def sub(m):
                s = m.group(1)
                r_ = s[1:] if s.startswith("@") else f"{s}{row}"
                v = cell(sheet, r_)
                if v is None or placeholder(v) or str(v).strip() == "":
                    raise KeyError(s)
                v = str(v).strip()
                return str(int(float(v))) if re.fullmatch(r"-?\d+\.0", v) else v
            try:
                out.append(re.sub(r"\{([^{}]+)\}", sub, tok))
            except KeyError:
                return None
        return out

    ap_ = Applier(L)
    changes, groups = [], {}
    for sheet, ref, path, conv, built in binds:
        if sheet not in wb.sheetnames:
            continue
        cur = cell(sheet, ref)
        b = json.loads(built) if built not in (None, "") else None
        toks = resolve(path, sheet, ref)
        if toks is None:
            continue
        gk = "/".join(toks[:max((i for i, t in enumerate(toks) if TOKEN.match(t).group(2) is not None), default=0) + 1])
        g_ = groups.setdefault(gk, {"toks": toks, "now": False, "was": False})
        if norm(cur) is not None and not placeholder(cur):
            g_["now"] = True
        if norm(b) is not None and not placeholder(b):
            g_["was"] = True
        if norm(cur) == norm(b):
            continue
        if placeholder(cur):
            continue
        new = convert(cur, conv or None)
        old = ap_.get(toks)
        if new == old:
            continue
        if ap_.set(toks, new, (sheet, coordinate_to_tuple(ref)[0])):
            changes.append({"sheet": sheet, "cell": ref, "field": "/".join(toks), "was": old, "now": new})

    removed = []
    for gk, g_ in groups.items():
        if g_["was"] and not g_["now"]:
            hit = ap_.item_of(g_["toks"])
            if hit:
                removed.append((gk, hit))
    for gk, (lst, it) in removed:
        if any(x is it for x in lst):
            lst.remove(next(x for x in lst if x is it))
    for lst, pad in ap_.padded:
        if not pad and any(x is pad for x in lst):
            lst.remove(next(x for x in lst if x is pad))

    # keep derived fields consistent
    gb = (((L.get("company_layer") or {}).get("internal") or {}).get("growth_barriers") or {})
    bind_ = [b_.get("barrier") for b_ in (gb.get("barriers") or []) if isinstance(b_, dict) and b_.get("status") == "binding"]
    if len(bind_) == 1 and gb.get("binding") != bind_[0]:
        gb["binding"] = bind_[0]
        changes.append({"sheet": "Growth Barriers", "cell": "B6:B11", "field": "company_layer/internal/growth_barriers/binding",
                        "was": None, "now": bind_[0]})

    added = len(ap_.new_rows) + sum(1 for lst, pad in ap_.padded if pad)
    by_sheet = {}
    for c in changes:
        by_sheet[c["sheet"]] = by_sheet.get(c["sheet"], 0) + 1
    report = {"workbook": os.path.basename(a.workbook), "cells_changed": len(changes), "rows_added": added,
              "rows_removed": len(removed), "by_sheet": by_sheet,
              "changes": [{k: v for k, v in c.items()} for c in changes]}
    if a.dry_run:
        print(json.dumps(report, indent=1, default=str))
        return
    if not changes and not removed:
        print(json.dumps({**report, "note": "No input cell differs from the ledger; nothing written."}, default=str))
        return
    hist = L.setdefault("workbook_sync", [])
    if isinstance(hist, dict):
        hist = L["workbook_sync"] = [hist]
    hist.append({"at": datetime.datetime.now().isoformat(timespec="seconds"), "workbook": os.path.basename(a.workbook),
                 "cells_changed": len(changes), "rows_added": added, "rows_removed": len(removed), "by_sheet": by_sheet})
    del hist[:-20]
    out = a.out or a.ledger
    if not out:
        out = os.path.splitext(a.workbook)[0] + "-ledger.json"
    if a.ledger and not a.out:
        shutil.copyfile(a.ledger, os.path.splitext(a.ledger)[0] + ".before-sync.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump(L, f, indent=2, ensure_ascii=False, default=str)
    print(json.dumps({**report, "out": out}, indent=1, default=str))


if __name__ == "__main__":
    main()
